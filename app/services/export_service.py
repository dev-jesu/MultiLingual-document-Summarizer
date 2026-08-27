import io
import datetime
from typing import Optional
import html

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.config import settings
from app.fonts.font_loader import get_registered_font_name
from app.utils.language_utils import get_language_name, normalize_language_code
from app.utils.logger import logger


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass ReportLab canvas to compute and render 'Page X of Y' dynamic footers.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Footer divider line
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.75)
        self.line(54, 48, self._pagesize[0] - 54, 48)

        # Left footer text
        left_footer = f"{settings.REPORT_TITLE} — Academic Project"
        self.drawString(54, 32, left_footer)

        # Right footer page numbers
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(self._pagesize[0] - 54, 32, page_str)
        self.restoreState()


class ExportService:
    """
    Service for generating professional downloadable reports in PDF, DOCX, and TXT formats.
    """

    @staticmethod
    def _format_language_label(language_code_or_name: Optional[str]) -> str:
        """Format a human-readable language string e.g. 'Tamil (தமிழ்)' or 'English'."""
        normalized = normalize_language_code(language_code_or_name)
        if not normalized:
            return language_code_or_name or "English"

        english_name = get_language_name(normalized, native=False)
        native_name = get_language_name(normalized, native=True)

        if english_name == native_name:
            return english_name
        return f"{english_name} ({native_name})"

    @staticmethod
    def export_pdf(
        filename: str,
        summary: str,
        language: Optional[str] = "en",
        processing_time: Optional[str] = "N/A",
    ) -> io.BytesIO:
        """
        Generates a professionally styled PDF report using ReportLab.
        """
        logger.info("Generating PDF report for %s [Format: PDF]", filename)
        buffer = io.BytesIO()

        font_name = get_registered_font_name()
        lang_display = ExportService._format_language_label(language)
        generated_date = datetime.datetime.now().strftime("%B %d, %Y at %I:%M %p")

        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=64,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=4,
            fontStyle="Bold"
        )

        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#2563EB"),
            spaceAfter=12
        )

        meta_label_style = ParagraphStyle(
            "MetaLabel",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#475569")
        )

        meta_value_style = ParagraphStyle(
            "MetaValue",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#0F172A")
        )

        heading_style = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontName=font_name,
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#1E293B"),
            spaceBefore=14,
            spaceAfter=8,
            keepWithNext=True
        )

        body_style = ParagraphStyle(
            "SummaryBody",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=10,
            leading=15,
            textColor=colors.HexColor("#334155"),
            spaceAfter=10
        )

        elements = []

        # Header Title
        elements.append(Paragraph(html.escape(settings.REPORT_TITLE), title_style))
        elements.append(Paragraph("Multilingual Legal Document Summarization Tool", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563EB"), spaceBefore=0, spaceAfter=14))

        # Metadata Section
        meta_data = [
            [
                Paragraph("<b>Original File:</b>", meta_label_style),
                Paragraph(html.escape(filename), meta_value_style),
                Paragraph("<b>Generated On:</b>", meta_label_style),
                Paragraph(html.escape(generated_date), meta_value_style),
            ],
            [
                Paragraph("<b>Language:</b>", meta_label_style),
                Paragraph(html.escape(lang_display), meta_value_style),
                Paragraph("<b>Processing Time:</b>", meta_label_style),
                Paragraph(html.escape(str(processing_time)), meta_value_style),
            ]
        ]

        meta_table = Table(meta_data, colWidths=[90, 160, 90, 164])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor("#E2E8F0")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#F1F5F9")),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))

        elements.append(meta_table)
        elements.append(Spacer(1, 16))

        # Summary Title & Content
        elements.append(Paragraph("Legal Summary", heading_style))
        elements.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor("#CBD5E1"), spaceBefore=2, spaceAfter=10))

        paragraphs = summary.strip().split("\n")
        for para in paragraphs:
            para_text = para.strip()
            if para_text:
                escaped_text = html.escape(para_text).replace("\t", "&nbsp;&nbsp;&nbsp;&nbsp;")
                elements.append(Paragraph(escaped_text, body_style))

        doc.build(elements, canvasmaker=NumberedCanvas)
        buffer.seek(0)
        logger.info("PDF report generation complete for %s", filename)
        return buffer

    @staticmethod
    def export_docx(
        filename: str,
        summary: str,
        language: Optional[str] = "en",
        processing_time: Optional[str] = "N/A",
    ) -> io.BytesIO:
        """
        Generates a professionally formatted DOCX report using python-docx.
        """
        logger.info("Generating DOCX report for %s [Format: DOCX]", filename)
        document = docx.Document()

        lang_display = ExportService._format_language_label(language)
        generated_date = datetime.datetime.now().strftime("%B %d, %Y at %I:%M %p")

        # Configure Margins
        sections = document.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Base Font Configuration
        normal_style = document.styles['Normal']
        font = normal_style.font
        font.name = 'Nirmala UI'
        font.size = Pt(10.5)
        font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        # Title
        p_title = document.add_paragraph()
        run_title = p_title.add_run(settings.REPORT_TITLE)
        run_title.font.name = 'Nirmala UI'
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p_title.paragraph_format.space_after = Pt(2)

        # Subtitle
        p_sub = document.add_paragraph()
        run_sub = p_sub.add_run("Multilingual Legal Document Summarization Tool")
        run_sub.font.name = 'Nirmala UI'
        run_sub.font.size = Pt(11)
        run_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        p_sub.paragraph_format.space_after = Pt(12)

        # Document Information Table
        table = document.add_table(rows=2, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers_values = [
            ("Original File:", filename, "Generated On:", generated_date),
            ("Language:", lang_display, "Processing Time:", str(processing_time))
        ]

        for r_idx, row_data in enumerate(headers_values):
            row = table.rows[r_idx]
            for c_idx, cell_text in enumerate(row_data):
                cell = row.cells[c_idx]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                run = p.add_run(cell_text)
                run.font.name = 'Nirmala UI'
                run.font.size = Pt(9)
                if c_idx % 2 == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
                else:
                    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        document.add_paragraph().paragraph_format.space_after = Pt(8)

        # Summary Heading
        p_head = document.add_paragraph()
        run_head = p_head.add_run("Legal Summary")
        run_head.font.name = 'Nirmala UI'
        run_head.font.size = Pt(14)
        run_head.font.bold = True
        run_head.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p_head.paragraph_format.space_before = Pt(12)
        p_head.paragraph_format.space_after = Pt(6)

        # Summary Paragraphs
        for para in summary.strip().split("\n"):
            para_text = para.strip()
            if para_text:
                p_body = document.add_paragraph()
                p_body.paragraph_format.space_after = Pt(8)
                p_body.paragraph_format.line_spacing = 1.2
                run_b = p_body.add_run(para_text)
                run_b.font.name = 'Nirmala UI'
                run_b.font.size = Pt(10.5)
                run_b.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        # Footer
        footer = document.sections[0].footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_foot = p_foot.add_run(f"Generated using {settings.REPORT_TITLE} | Academic Project")
        run_foot.font.name = 'Nirmala UI'
        run_foot.font.size = Pt(8.5)
        run_foot.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        logger.info("DOCX report generation complete for %s", filename)
        return buffer

    @staticmethod
    def export_txt(
        filename: str,
        summary: str,
        language: Optional[str] = "en",
        processing_time: Optional[str] = "N/A",
    ) -> bytes:
        """
        Generates a clean UTF-8 plain text report string.
        """
        logger.info("Generating TXT report for %s [Format: TXT]", filename)
        lang_display = ExportService._format_language_label(language)
        generated_date = datetime.datetime.now().strftime("%B %d, %Y at %I:%M %p")

        divider = "=" * 70
        subdivider = "-" * 70

        lines = [
            divider,
            f"{settings.REPORT_TITLE}",
            "Multilingual Legal Document Summarization Tool",
            divider,
            "",
            "DOCUMENT INFORMATION",
            subdivider,
            f"Original File   : {filename}",
            f"Language        : {lang_display}",
            f"Generated On    : {generated_date}",
            f"Processing Time : {processing_time}",
            subdivider,
            "",
            "LEGAL SUMMARY",
            subdivider,
            summary.strip(),
            "",
            divider,
            f"Generated using {settings.REPORT_TITLE} | Academic Project",
            divider,
        ]

        text_content = "\n".join(lines)
        return text_content.encode("utf-8")
